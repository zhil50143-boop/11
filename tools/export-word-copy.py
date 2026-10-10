"""Export the complete original copy into one editable, source-labelled Word file."""
from pathlib import Path
from collections import defaultdict, Counter
from datetime import datetime
from zoneinfo import ZoneInfo
import hashlib
import json
import re
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
TARGET = Path(sys.argv[1]).resolve()
RECORD = Path(sys.argv[2]).resolve()
read = lambda p: json.loads(Path(p).read_text(encoding="utf-8-sig"))
digest = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
index = read(ROOT / "docs/PLATFORM_REVIEW_COPY.index.json")
for source in index["inputs"]:
    text = (ROOT / source["file"]).read_text(encoding="utf-8").replace("\r\n", "\n")
    assert hashlib.sha256(text.encode()).hexdigest() == source["sha256"], source["file"]
assert digest(ROOT / "docs/PLATFORM_REVIEW_COPY.md") == index["export"]["sha256"]
assert not TARGET.exists(), "Preserve prior deliverables"

# These four player-facing strings do not contain Chinese and were omitted by
# the conservative Chinese-string inventory. Keep each distinct source record.
json_extras = [
    {"file": "assets/resources/data/story/chapter01/chapter01_manifest.json", "path": "episodes[3].name", "text": "01.mp3"},
    {"file": "assets/resources/data/story/chapter01/ep03_dinner.json", "path": "nodes[4].title", "text": "23:18"},
    {"file": "assets/resources/data/story/chapter01/ep04_mp3.json", "path": "name", "text": "01.mp3"},
    {"file": "assets/resources/data/story/chapter01/ep04_mp3.json", "path": "nodes[1].text", "text": "01.mp3"},
]
def field_at(obj, pointer):
    for token in re.findall(r"[^.\[\]]+", pointer):
        obj = obj[int(token)] if isinstance(obj, list) else obj[token]
    return obj
for item in json_extras:
    assert field_at(read(ROOT / item["file"]), item["path"]) == item["text"]
ui_extras = [{"file": "assets/scripts/ui/VisualDraft.ts", "line": n, "text": value,
              "format": "literal", "scope": "design-preview-only"}
             for n, value in [(33, "01.mp3"), (35, "Aa"), (44, "Aa")]]
for item in ui_extras:
    assert item["text"] in (ROOT / item["file"]).read_text(encoding="utf-8").splitlines()[item["line"] - 1]

json_records = index["textRecords"] + json_extras
ui_records = sorted(index["interfaceText"] + ui_extras, key=lambda item: (item["file"], item["line"]))
audio = read(ROOT / "art/original/audio/manifest.json")
store_snapshot = read(ROOT / "docs/PLATFORM_REVIEW_COPY.store.json")
assert store_snapshot["appId"] == 966933
store_modules = store_snapshot["modules"]
store = [("游戏名称", "basic-info", "title"), ("游戏简介", "basic-info", "description"),
         ("开发者的话", "developer-info", "developer_message"), ("推荐语", "profile-promotion", "promotion_text")]
assert store_modules["basic-info"]["title"]["current_value"] == "遗憾"
assert (len(json_records), len(ui_records), len(audio["events"]), len(store)) == (3303, 175, 15, 4)

document = Document()
section = document.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.7)
section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
section.footer_distance = Cm(0.7)

def font(style, size, bold=False):
    style.font.name = "SimSun"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
        fonts.set(qn("w:" + key), "SimSun")
    for key in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        fonts.attrib.pop(qn("w:" + key), None)

for name, size, bold in [("Normal", 10.5, False), ("Title", 24, True),
                         ("Subtitle", 12, False), ("Heading 1", 16, True),
                         ("Heading 2", 12.5, True), ("Heading 3", 11, True)]:
    font(document.styles[name], size, bold)
for border in list(document.styles.element.iter(qn("w:pBdr"))):
    border.getparent().remove(border)
document.styles["Subtitle"].font.italic = False
normal = document.styles["Normal"].paragraph_format
normal.line_spacing = 1.2
normal.space_after = Pt(3)
normal.widow_control = True
for name in ["Heading 1", "Heading 2", "Heading 3"]:
    pf = document.styles[name].paragraph_format
    pf.keep_with_next = True
    pf.space_before, pf.space_after = Pt(10), Pt(5)
    pf.line_spacing = 1.15
document.styles["Heading 1"].paragraph_format.page_break_before = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("第 ").font.size = Pt(9)
for token, text in [("begin", None), (None, " PAGE "), ("end", None)]:
    run = footer.add_run()
    element = OxmlElement("w:fldChar" if token else "w:instrText")
    if token:
        element.set(qn("w:fldCharType"), token)
    else:
        element.set(qn("xml:space"), "preserve")
        element.text = text
    run._r.append(element)
footer.add_run(" 页").font.size = Pt(9)

counts = Counter()
coverage = []
bookmark_id = 1
def bookmark(paragraph, name):
    global bookmark_id
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(bookmark_id)); start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)
    bookmark_id += 1

def metadata(text, keep=True):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(80, 80, 80)
    return p

def heading(text, level=1, anchor=None):
    p = document.add_heading(text, level)
    if anchor:
        bookmark(p, anchor)
    return p

def record(item, kind, label, paragraph=None):
    p = paragraph if paragraph is not None else document.add_paragraph()
    if paragraph is not None and p.text:
        p.add_run("   ")
    run = p.add_run(label + "  ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(80, 80, 80)
    p.add_run(item["text"])
    counts[kind] += 1
    tag = f"text_{len(coverage) + 1:05d}"
    bookmark(p, tag)
    coverage.append({"kind": kind, "bookmark": tag, **item})
    return p

stamp = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y年%m月%d日")
document.add_paragraph("遗憾 全部文本汇总", "Title")
document.add_paragraph("游戏全分支与后台文案审核用", "Subtitle")
document.add_paragraph(f"版本 1.0.0-rc.3    汇总日期 {stamp}")
document.add_paragraph("本文件完整收录十章所有分支、选项、旧物调查、照片背文、信件、短信、录音、六个结局，以及界面提示和后台商店文案。所有文本保留原句，重复原文按出处分别收录，包含剧透和隐藏结局。")
document.add_paragraph("互斥选项与条件分支并列收录，不能按文档顺序理解为同一路线。每段附节点编号或来源字段，便于核对；独立设计预览和内部背景文字另行标明。")
document.add_paragraph("收录范围 10章 60片段 467节点 6结局")
document.add_paragraph("来源条目 3303条JSON文字 175条源码文字 15条录音口述 4个后台字段 共3497条")
metadata("剧情与界面源码基线 " + index["sourceCommit"], False)
metadata("原项目名 余生未寄    仓库 https://github.com/zhil50143-boop/11", False)

document.add_page_break()
toc_heading = document.add_heading("目录", 1)
toc_heading.paragraph_format.page_break_before = False
catalog = read(ROOT / "assets/resources/data/story/catalog.json")
chapter_info = []
for i, ref in enumerate(catalog["chapters"], 1):
    file = "assets/resources/" + ref["resource"] + ".json"
    chapter_info.append((i, file, read(ROOT / file)))
sections = [("后台商店文案", "store"), ("人物显示名称与默认文字", "common")]
sections += [(f"第{i}章 {data['title']}", f"chapter_{i}") for i, _, data in chapter_info]
sections += [("界面操作与异常提示", "ui"), ("独立设计预览文字", "preview"), ("实际录音口述台词", "voice")]
for label, anchor in sections:
    p = document.add_paragraph()
    link = OxmlElement("w:hyperlink"); link.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = label
    r.append(t); link.append(r); p._p.append(link)

heading("后台商店文案", anchor="store")
metadata("深海游戏工作室 473536    新游戏 遗憾 966933    本次汇总前重新读取后台", False)
for label, module, field in store:
    heading(label, 2)
    record({"file": f"TapTap/{module}", "path": field, "text": store_modules[module][field]["current_value"]}, "store", "原文")

by_file = defaultdict(list)
for item in json_records:
    by_file[item["file"]].append(item)
seen = set()
heading("人物显示名称与默认文字", anchor="common")
for file in ["assets/resources/data/presentation.json", "assets/resources/data/story/catalog.json"]:
    if not by_file[file]:
        continue
    heading("人物与展示设置" if file.endswith("presentation.json") else "章节目录文字", 2)
    metadata("来源 " + file)
    for item in by_file[file]:
        record(item, "json", item["path"])
        seen.add((item["file"], item["path"]))

type_names = {"passage":"生活正文", "phone":"电话与短信", "dialogue":"对白", "narration":"叙述", "choice":"重要决定", "investigation":"旧物调查", "photo":"照片", "letter":"信件与纸页", "audioInteraction":"录音", "transition":"时间转场", "ending":"结局", "condition":"条件节点", "endingRoute":"结局路由", "episodeEnd":"片段结尾"}
field_labels = {"title":"标题", "text":"正文", "backText":"照片背面", "transcript":"录音文字", "prompt":"选择提示", "consequence":"背景说明", "actionText":"操作文字", "requiredHint":"调查提示", "doneText":"收起操作", "from":"转场起点", "to":"转场终点"}
for number, manifest_file, manifest in chapter_info:
    heading(f"第{number}章 {manifest['title']}", anchor=f"chapter_{number}")
    metadata("章节来源 " + manifest_file)
    for item in by_file[manifest_file]:
        record(item, "json", item["path"]).paragraph_format.keep_with_next = True
        seen.add((item["file"], item["path"]))
    for ep in manifest["episodes"]:
        episode_file = "assets/resources/" + ep["resource"] + ".json"
        episode = read(ROOT / episode_file)
        heading(episode["name"], 2)
        metadata("片段来源 " + episode_file)
        items = by_file[episode_file]
        for item in items:
            if not item["path"].startswith("nodes["):
                record(item, "json", item["path"]).paragraph_format.keep_with_next = True
                seen.add((item["file"], item["path"]))
        for n, node in enumerate(episode["nodes"]):
            node_items = [item for item in items if item["path"].startswith(f"nodes[{n}].")]
            if not node_items:
                continue
            p = metadata(node["id"] + "    " + type_names.get(node["type"], node["type"]))
            p.paragraph_format.space_before = Pt(6)
            grouped = document.add_paragraph()
            grouped.paragraph_format.space_after = Pt(4)
            grouped.paragraph_format.keep_with_next = any(
                item["path"].split("].", 1)[1] not in ["title", "lifeContext.time.label", "lifeContext.time.location"]
                for item in node_items)
            for item in node_items:
                relative = item["path"].split("].", 1)[1]
                if relative in ["title", "lifeContext.time.label", "lifeContext.time.location"]:
                    label = {"title":"标题", "lifeContext.time.label":"时间", "lifeContext.time.location":"地点"}[relative]
                    record(item, "json", label, grouped)
                else:
                    if relative.startswith("paragraphs["):
                        part = int(re.search(r"\[(\d+)\]", relative).group(1))
                        speaker = episode["nodes"][n]["paragraphs"][part].get("speaker")
                        display = read(ROOT / "assets/resources/data/presentation.json")["speakers"].get(speaker)
                        label = display or "叙述"
                    elif relative.startswith("options["):
                        label = "选项 " + str(int(re.search(r"\[(\d+)\]", relative).group(1)) + 1)
                    elif relative.startswith("items["):
                        label = "调查入口 " + str(int(re.search(r"\[(\d+)\]", relative).group(1)) + 1)
                    else:
                        label = field_labels.get(relative, relative)
                    field_paragraph = record(item, "json", label)
                    if relative == "actionText":
                        field_paragraph.paragraph_format.keep_with_next = True
                seen.add((item["file"], item["path"]))
            if not grouped.text:
                grouped._element.getparent().remove(grouped._element)

assert len(seen) == len(json_records), (len(seen), len(json_records))
for scope, title, anchor in [("game-source","界面操作与异常提示","ui"), ("design-preview-only","独立设计预览文字","preview")]:
    heading(title, anchor=anchor)
    if scope == "design-preview-only":
        document.add_paragraph("本节为独立布局预览中的文字，按源码完整保留，不作为正常剧情界面已展示的证明。")
    else:
        document.add_paragraph("包括按钮、设置、阅读、翻面、播放、重读、存档、加载及异常提示。动态模板保留源码写法；内部诊断和背景常量也按来源收录。")
    file = None
    for item in ui_records:
        if item["scope"] != scope:
            continue
        if item["file"] != file:
            file = item["file"]
            metadata("来源 " + file)
        record(item, "source", "第" + str(item["line"]) + "行")

heading("实际录音口述台词", anchor="voice")
document.add_paragraph("以下按实际音频制作清单收录第三版两名角色的口述台词。完整录音约63.53秒，第一章使用前16.66秒，第八章使用完整文件；音色保持用户确认的版本。")
metadata("来源 art/original/audio/manifest.json")
for event_number, event in enumerate(audio["events"]):
    record({"file":"art/original/audio/manifest.json", "path":f"events[{event_number}].text", "text":event["text"]}, "voice", f"{event['role']}  {event['start']:.2f}至{event['end']:.2f}秒")

assert counts == {"json": 3303, "source": 175, "voice": 15, "store": 4}, counts
TARGET.parent.mkdir(parents=True, exist_ok=True)
document.core_properties.title = "遗憾 全部文本汇总"
document.core_properties.subject = "十章全分支与后台文案审核"
document.core_properties.author = "深海游戏工作室"
document.core_properties.comments = ""
document.save(TARGET)

# Verify editable OOXML text, not just file creation or ZIP validity.
loaded = Document(TARGET)
all_text = "\n".join(p.text for p in loaded.paragraphs)
for item in coverage:
    assert item["text"] in all_text, (item.get("file"), item.get("path"))
names = {item.get(qn("w:name")) for item in loaded.element.iter(qn("w:bookmarkStart"))}
assert all(item["bookmark"] in names for item in coverage)
RECORD.parent.mkdir(parents=True, exist_ok=True)
write_record = {"file":str(TARGET), "sha256":digest(TARGET), "bytes":TARGET.stat().st_size, "generatedAt":datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(), "sourceCommit":index["sourceCommit"], "sourceInputsVerified":len(index["inputs"]), "counts":dict(counts), "totalSourceEntries":len(coverage), "missingTextEntries":0, "sourceMarkdownSha256":index["export"]["sha256"], "coverage":coverage}
RECORD.write_text(json.dumps(write_record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({key:write_record[key] for key in ["file","bytes","sha256","counts","totalSourceEntries","missingTextEntries"]}, ensure_ascii=False, indent=2))
